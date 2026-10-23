"""Testes das funções de edição das planilhas (planilhas.linhas_* e planilhas.salvar_*), em arquivos temporários."""

import openpyxl
import pytest

from automacao import planilhas


def criar(caminho, abas: dict[str, list]) -> None:
    livro = openpyxl.Workbook()
    livro.remove(livro.active)
    for nome, linhas in abas.items():
        folha = livro.create_sheet(nome)
        for linha in linhas:
            folha.append(linha if isinstance(linha, list) else [linha])
    livro.save(caminho)


def conteudo(caminho) -> dict[str, list]:
    livro = openpyxl.load_workbook(caminho)
    try:
        return {folha.title: [linha[0] for linha in folha.iter_rows(values_only=True)]
                for folha in livro.worksheets}
    finally:
        livro.close()


# ------------------------------------------------------------ leitura

def test_planilha_que_nao_existe_vem_vazia(tmp_path):
    assert planilhas.linhas_perfis(tmp_path / "nao.xlsx") == {"instagram": [], "tiktok": []}
    assert planilhas.linhas_destinatarios(tmp_path / "nao.xlsx") == {"instagram": [], "tiktok": []}


def test_le_as_abas_sem_o_cabecalho_e_com_linhas_invalidas(tmp_path):
    caminho = tmp_path / "pa.xlsx"
    criar(caminho, {
        "instagram": ["url", "instagram.com/fulano", None, "  https://www.instagram.com/ciclano/  ", "lixo"],
        "TikTok": ["https://www.tiktok.com/@primeira"],   # sem cabeçalho e com maiúscula no nome
        "outra": ["https://www.instagram.com/ignorado/"],
    })
    assert planilhas.linhas_perfis(caminho) == {
        "instagram": ["instagram.com/fulano", "https://www.instagram.com/ciclano/", "lixo"],
        "tiktok": ["https://www.tiktok.com/@primeira"],
    }


def test_le_destinatarios_reconhece_o_cabecalho(tmp_path):
    caminho = tmp_path / "pc.xlsx"
    criar(caminho, {"instagram": ["perfil (somente o @)", "@fulano", "ciclano"],
                    "tiktok": ["beltrano"]})
    assert planilhas.linhas_destinatarios(caminho) == {"instagram": ["@fulano", "ciclano"],
                                                       "tiktok": ["beltrano"]}


def test_caminho_padrao_e_lido_na_hora(tmp_path, monkeypatch):
    caminho = tmp_path / "padrao.xlsx"
    criar(caminho, {"instagram": ["url", "instagram.com/fulano"]})
    monkeypatch.setattr(planilhas, "PLANILHA_PERFIS", caminho)
    assert planilhas.linhas_perfis()["instagram"] == ["instagram.com/fulano"]


# ------------------------------------------------------------ gravação

def test_salvar_cria_a_planilha_com_abas_e_cabecalhos(tmp_path):
    caminho = tmp_path / "data" / "pa.xlsx"   # a pasta também é criada
    planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/fulano/"], "tiktok": []}, caminho)
    assert conteudo(caminho) == {"instagram": ["url", "https://www.instagram.com/fulano/"],
                                 "tiktok": ["url"]}
    assert not list(caminho.parent.glob(".*salvando*"))  # o temporário some


def test_salvar_destinatarios_tira_o_arroba(tmp_path):
    caminho = tmp_path / "pc.xlsx"
    planilhas.salvar_destinatarios({"instagram": ["@fulano", " ciclano "], "tiktok": ["@beltrano"]}, caminho)
    assert conteudo(caminho) == {"instagram": ["perfil (somente o @)", "fulano", "ciclano"],
                                 "tiktok": ["perfil (somente o @)", "beltrano"]}


def test_salvar_por_cima_mantem_nome_das_abas_cabecalho_e_outras_abas(tmp_path):
    caminho = tmp_path / "pa.xlsx"
    criar(caminho, {"Instagram": ["Links", "https://www.instagram.com/velho/", "https://www.instagram.com/outro/"],
                    "notas": ["não mexer"]})
    planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/novo/"],
                             "tiktok": ["https://www.tiktok.com/@novo"]}, caminho)
    assert conteudo(caminho) == {
        "Instagram": ["Links", "https://www.instagram.com/novo/"],
        "notas": ["não mexer"],
        "tiktok": ["url", "https://www.tiktok.com/@novo"],
    }


def test_salvar_mantem_outras_colunas_junto_do_seu_perfil(tmp_path):
    """Anotações em outras colunas não podem ser apagadas nem trocar de perfil."""
    caminho = tmp_path / "pa.xlsx"
    criar(caminho, {"instagram": [["url", "Anotação"],
                                  ["instagram.com/Fulano", "cliente VIP"],
                                  ["https://www.instagram.com/ciclano/", "não responder"],
                                  ["https://www.instagram.com/sai/", "vai embora"]]})
    planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/ciclano/",
                                           "https://www.instagram.com/beltrano/",
                                           "https://www.instagram.com/fulano/"]}, caminho)
    livro = openpyxl.load_workbook(caminho)
    linhas = [linha for linha in livro["instagram"].iter_rows(values_only=True) if any(linha)]
    livro.close()
    assert linhas == [
        ("url", "Anotação"),
        ("https://www.instagram.com/ciclano/", "não responder"),
        ("https://www.instagram.com/beltrano/", None),
        ("https://www.instagram.com/fulano/", "cliente VIP"),
    ]


def test_salvar_aba_sem_cabecalho_ganha_cabecalho(tmp_path):
    caminho = tmp_path / "pc.xlsx"
    criar(caminho, {"instagram": ["fulano", "ciclano"], "tiktok": ["perfil (somente o @)"]})
    planilhas.salvar_destinatarios({"instagram": ["fulano"]}, caminho)
    dados = conteudo(caminho)
    assert dados["instagram"] == ["perfil (somente o @)", "fulano"]
    assert dados["tiktok"] == ["perfil (somente o @)"]  # aba não pedida fica como estava


def test_salvar_lista_vazia_deixa_so_o_cabecalho(tmp_path):
    caminho = tmp_path / "pa.xlsx"
    planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/a/"]}, caminho)
    planilhas.salvar_perfis({"instagram": []}, caminho)
    assert conteudo(caminho)["instagram"] == ["url"]


def test_ida_e_volta_com_os_leitores_do_interagir(tmp_path):
    """O que a interface grava é exatamente o que ler_perfis e ler_destinatarios usam."""
    perfis = tmp_path / "pa.xlsx"
    destinatarios = tmp_path / "pc.xlsx"
    planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/fulano/"],
                             "tiktok": ["https://www.tiktok.com/@ciclano"]}, perfis)
    planilhas.salvar_destinatarios({"instagram": ["@amigo"], "tiktok": ["outro.amigo"]}, destinatarios)

    assert planilhas.ler_perfis(perfis) == ["https://www.instagram.com/fulano/", "https://www.tiktok.com/@ciclano"]
    assert planilhas.ler_destinatarios(destinatarios) == {"Instagram": ["amigo"], "TikTok": ["outro.amigo"]}
    assert planilhas.linhas_perfis(perfis) == {"instagram": ["https://www.instagram.com/fulano/"],
                                               "tiktok": ["https://www.tiktok.com/@ciclano"]}


def test_erro_ao_gravar_nao_estraga_a_planilha(tmp_path, monkeypatch):
    caminho = tmp_path / "pa.xlsx"
    planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/original/"]}, caminho)

    def falha(origem, destino):
        raise PermissionError("arquivo aberto em outro programa")

    monkeypatch.setattr(planilhas.os, "replace", falha)
    with pytest.raises(PermissionError):
        planilhas.salvar_perfis({"instagram": ["https://www.instagram.com/novo/"]}, caminho)
    assert conteudo(caminho)["instagram"] == ["url", "https://www.instagram.com/original/"]
    assert not list(tmp_path.glob(".*salvando*"))


def test_funcoes_antigas_continuam_iguais(tmp_path, capsys):
    caminho = tmp_path / "pa.xlsx"
    criar(caminho, {"instagram": ["url", "instagram.com/fulano", "lixo"]})
    assert planilhas.ler_perfis(caminho) == ["instagram.com/fulano"]
    assert "linha 3" in capsys.readouterr().out
